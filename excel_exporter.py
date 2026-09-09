"""
엑셀 파일 저장 모듈 (excel_exporter.py)
pandas와 openpyxl을 사용하여 수집된 음식점 및 메뉴 데이터를
지정된 형식(| 음식점명 | 메뉴명 |)의 엑셀 파일(restaurants_menu.xlsx)로 저장합니다.
"""

from pathlib import Path
from typing import List, Dict, Union
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """수집된 데이터를 엑셀(Excel) 파일로 저장하고 스타일을 서식화하는 클래스"""

    COLUMNS = ["음식점명", "메뉴명"]

    @classmethod
    def save_to_excel(
        cls,
        records: List[Dict[str, str]],
        output_path: Union[str, Path]
    ) -> Path:
        """
        데이터 레코드 리스트를 엑셀 파일로 저장합니다.

        :param records: [{'음식점명': '...', '메뉴명': '...'}, ...] 형태의 리스트
        :param output_path: 저장할 엑셀 파일의 경로
        :return: 저장된 파일의 Path 객체
        """
        output_path = Path(output_path)

        # 1. pandas DataFrame 생성
        if not records:
            df = pd.DataFrame(columns=cls.COLUMNS)
        else:
            df = pd.DataFrame(records)
            # 요구사항 컬럼 순서 보장
            df = df[cls.COLUMNS]

        # 2. openpyxl 엔진을 사용하여 엑셀 파일로 쓰기
        try:
            with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="음식점_메뉴_목록")
                workbook = writer.book
                worksheet = writer.sheets["음식점_메뉴_목록"]

                # 3. 스타일 서식 적용 (헤더 스타일, 열 너비 자동 맞춤, 테두리)
                cls._apply_excel_styles(worksheet, df)

            print(f"\n[성공] 엑셀 파일이 성공적으로 저장되었습니다: {output_path.resolve()}")
            return output_path

        except PermissionError:
            raise PermissionError(
                f"파일 '{output_path.name}'이(가) 다른 프로그램(예: Excel)에서 열려 있어 저장할 수 없습니다.\n"
                "파일을 닫은 후 다시 실행해주세요."
            )
        except Exception as e:
            raise RuntimeError(f"엑셀 저장 중 오류가 발생했습니다: {e}")

    @staticmethod
    def _apply_excel_styles(worksheet, df: pd.DataFrame):
        """엑셀 워크시트에 가독성을 높이는 서식 스타일을 적용합니다."""
        # 헤더 스타일
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")  # 짙은 블루
        header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="맑은 고딕", size=10)

        # 테두리 스타일
        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3")
        )

        # 헤더 행 서식 적용
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        worksheet.row_dimensions[1].height = 26

        # 데이터 셀 서식 적용
        for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=len(df) + 1), start=2):
            worksheet.row_dimensions[row_idx].height = 20
            for cell in row:
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="left", vertical="center")

        # 열 너비 자동 조절 (한글/영문 글자 길이에 맞춰 여유 있게 설정)
        for col_idx, col in enumerate(worksheet.columns, start=1):
            max_len = 0
            for cell in col:
                val = str(cell.value or "")
                # 한글은 대략 1.7배 폭 계산
                length = sum(1.7 if ord(c) > 127 else 1.0 for c in val)
                if length > max_len:
                    max_len = length
            col_letter = get_column_letter(col_idx)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 15)
